import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from pptx import Presentation
from pptx.dml.color import RGBColor

def get_fill(shape):
    try:
        f = shape.fill
        if f.type is None:
            return "no fill"
        if f.type.name == "SOLID":
            try:
                return f"solid #{f.fore_color.rgb}"
            except:
                return "solid (theme color)"
        return f.type.name
    except:
        return "?"

def get_font_color(para):
    for run in para.runs:
        try:
            c = run.font.color.rgb
            return f"#{c}"
        except:
            pass
    return "?"

prs = Presentation(r'G:\내 드라이브\Projects\Project-Archive\projects\솔타메디칼\docs\솔타메디칼_현업인터뷰결과보고서_V0.1_261013.pptx')

for i, slide in enumerate(prs.slides):
    print(f'\n=== Slide {i+1} ===')
    for sh in slide.shapes:
        fill = get_fill(sh)
        if fill not in ("no fill", "?", "BACKGROUND"):
            l = round(sh.left/914400*2.54,2)
            t = round(sh.top/914400*2.54,2)
            txt = sh.text_frame.text[:40].replace('\n','|') if sh.has_text_frame else ''
            print(f'  [{sh.name}] fill={fill} ({l},{t}) text="{txt}"')
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    fc = get_font_color(para)
                    if fc != "?":
                        pt = para.text[:30]
                        print(f'    font={fc} "{pt}"')
