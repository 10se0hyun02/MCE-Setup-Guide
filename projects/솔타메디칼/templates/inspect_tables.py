import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from pptx import Presentation

prs = Presentation(r'G:\내 드라이브\Projects\Project-Archive\projects\솔타메디칼\docs\솔타메디칼_현업인터뷰결과보고서_V0.1_261013.pptx')

for i, slide in enumerate(prs.slides):
    for sh in slide.shapes:
        if sh.shape_type == 19:  # TABLE
            print(f'\n=== Slide {i+1} - {sh.name} ===')
            tbl = sh.table
            for r in range(len(tbl.rows)):
                for c in range(len(tbl.columns)):
                    cell = tbl.cell(r, c)
                    print(f'  ({r},{c}): {repr(cell.text)}')
