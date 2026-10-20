import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from pptx import Presentation
from lxml import etree

prs = Presentation(r'G:\내 드라이브\Projects\Project-Archive\projects\솔타메디칼\docs\솔타메디칼_현업인터뷰결과보고서_V0.1_261013.pptx')

# 슬라이드 2의 주요 shape XML 확인
slide2 = prs.slides[1]
for sh in slide2.shapes:
    if sh.name in ('Shape 2', 'Shape 3', 'Text 4', 'Text 5', 'Text 6'):
        print(f'\n--- {sh.name} ---')
        print(etree.tostring(sh._element, pretty_print=True).decode('utf-8')[:800])

# 슬라이드 4 하단 노트 박스
slide4 = prs.slides[3]
for sh in slide4.shapes:
    if sh.name in ('Shape 30', 'Shape 2', 'Shape 4'):
        print(f'\n--- slide4/{sh.name} ---')
        print(etree.tostring(sh._element, pretty_print=True).decode('utf-8')[:600])

# 슬라이드 8 카드
slide8 = prs.slides[7]
for sh in slide8.shapes:
    if sh.name in ('Shape 2', 'Shape 3', 'Text 4'):
        print(f'\n--- slide8/{sh.name} ---')
        print(etree.tostring(sh._element, pretty_print=True).decode('utf-8')[:600])
