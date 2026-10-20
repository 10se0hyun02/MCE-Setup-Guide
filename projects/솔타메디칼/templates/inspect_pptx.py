import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from pptx import Presentation

prs = Presentation(r'G:\내 드라이브\Projects\Project-Archive\projects\솔타메디칼\docs\솔타메디칼_현업인터뷰결과보고서_V0.1_261013.pptx')
for i, slide in enumerate(prs.slides):
    print(f'\n=== Slide {i+1} [{slide.slide_layout.name}] ===')
    for sh in slide.shapes:
        l = round(sh.left/914400*2.54, 2)
        t = round(sh.top/914400*2.54, 2)
        w = round(sh.width/914400*2.54, 2)
        h = round(sh.height/914400*2.54, 2)
        st = sh.shape_type
        name = sh.name
        txt = ''
        if sh.has_text_frame:
            txt = sh.text_frame.text[:80].replace('\n', '|')
        elif st == 19:
            rows = len(sh.table.rows)
            cols = len(sh.table.columns)
            txt = f'[TABLE {rows}r x {cols}c]'
            for r in range(rows):
                for c in range(cols):
                    cell_txt = sh.table.cell(r, c).text[:30].replace('\n','|')
                    if cell_txt:
                        txt += f' ({r},{c}):{cell_txt}'
        print(f'  [{name}] type={st} ({l},{t}) {w}x{h}  "{txt}"')
